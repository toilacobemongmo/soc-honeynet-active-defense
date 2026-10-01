"""
Test proper OOXML for python-docx
"""

import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()
tbl = doc.add_table(rows=2, cols=2)
tbl.style = 'Table Grid'
cell = tbl.cell(0, 0)
shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="1F497D"/>')
cell._tc.get_or_add_tcPr().append(shd)
doc.save('./reports/test_valid.docx')
print("Test docx saved successfully.")
